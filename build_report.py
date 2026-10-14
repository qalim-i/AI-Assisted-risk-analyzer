from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
PEER_PDF = Path(r"C:\Users\Lenovo\Downloads\DSP_Report_Final_Manju.pdf")
OUT = ROOT / "output" / "AI_Assisted_Cyber_Risk_Management_Report.docx"
ASSETS = ROOT / "tmp" / "report_assets"
BLUE = "17365D"
BROWN = "8B3F0A"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:" + key), str(value))


def set_table_borders(table, color="D9D9D9", size="6"):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"val": "single", "sz": size, "color": color},
                            bottom={"val": "single", "sz": size, "color": color},
                            left={"val": "single", "sz": size, "color": color},
                            right={"val": "single", "sz": size, "color": color})


def set_page_border(section):
    sect_pr = section._sectPr
    borders = sect_pr.first_child_found_in("w:pgBorders")
    if borders is None:
        borders = OxmlElement("w:pgBorders")
        sect_pr.append(borders)
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "double")
        element.set(qn("w:sz"), "16")
        element.set(qn("w:space"), "20")
        element.set(qn("w:color"), "000000")
        borders.append(element)


def set_margins(section, top=0.55, bottom=0.55, left=0.85, right=0.85):
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)


def set_run_font(run, size=12, bold=False, color=BLACK, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(run, field):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def para(doc, text="", size=12, bold=False, align=None, space_after=7, indent=0, line=1.35, color=BLACK, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line
    if indent:
        fmt.first_line_indent = Inches(indent)
    r = p.add_run(text)
    set_run_font(r, size, bold, color, italic)
    return p


def heading(doc, text, level=1, centered=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(8)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 14 if level == 1 else 12.5, True)
    return p


def header_footer(section, title="AI ASSISTED CYBER RISK ANALYZER", page_number=False):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    h = section.header.paragraphs[0]
    h.clear()
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.border_bottom = None if False else None
    r = h.add_run(title)
    set_run_font(r, 9.5)
    p_pr = h._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), BROWN)
    pbdr.append(bottom)
    p_pr.append(pbdr)
    f = section.footer.paragraphs[0]
    f.clear()
    f.paragraph_format.space_before = Pt(2)
    tab_stops = f.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(3.2))
    tab_stops.add_tab_stop(Inches(6.35))
    rf = f.add_run("Dept. of AIML\tAY 2026-27\t")
    set_run_font(rf, 9.5, True)
    if page_number:
        rp = f.add_run()
        set_run_font(rp, 9.5, True)
        add_field(rp, "PAGE")
    p_pr = f._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "18")
    top.set(qn("w:space"), "5")
    top.set(qn("w:color"), BROWN)
    pbdr.append(top)
    p_pr.append(pbdr)


def clear_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].clear()
    section.footer.paragraphs[0].clear()


def make_logos():
    ASSETS.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(str(PEER_PDF))
    images = list(reader.pages[0].images)
    paths = []
    for i, image in enumerate(images[:2]):
        ext = Path(image.name).suffix or ".jpg"
        output = ASSETS / f"logo_{i + 1}{ext}"
        output.write_bytes(image.data)
        paths.append(output)
    return paths


def image_cell(cell, path, width):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))


def cover_page(doc, logos):
    s = doc.sections[-1]
    set_margins(s, 0.62, 0.55, 0.65, 0.65)
    set_page_border(s)
    clear_header_footer(s)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(4.65)
    table.columns[2].width = Inches(1.25)
    image_cell(table.cell(0, 0), logos[0], 0.78)
    image_cell(table.cell(0, 2), logos[1], 0.78)
    c = table.cell(0, 1)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, size in [("DAYANANDA SAGAR COLLEGE OF", 17), ("ENGINEERING", 17)]:
        r = p.add_run(text + "\n")
        set_run_font(r, size, True)
    for text, size in [("(An Autonomous Institute affiliated to Visvesvaraya Technological University (VTU),", 8.5),
                       ("Belagavi,", 8.5), ("Approved by AICTE and UGC, Accredited by NAAC with ‘A’ grade & ISO 9001 - 2015", 8.5),
                       ("Certified Institution)", 8.5), ("Shavige Malleshwara Hills, Kumaraswamy Layout, Bengaluru-560 111,", 9.5), ("India", 9.5)]:
        r = p.add_run(text + "\n")
        set_run_font(r, size)
    p.paragraph_format.space_after = Pt(0)
    para(doc, "", space_after=17)
    para(doc, "DEPARTMENT OF ARTIFICIAL INTELLIGENCE & MACHINE LEARNING", 12.5, True, WD_ALIGN_PARAGRAPH.CENTER, 20)
    para(doc, "Alternate Assessment Tool", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 34)
    para(doc, "Micro Project Report", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 26, color=BLUE)
    para(doc, "AI ASSISTED ENTERPRISE CYBER RISK MANAGEMENT", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 10, color=BLUE)
    para(doc, "USING LARGE LANGUAGE MODELS", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 30, color=BLUE)
    para(doc, "Submitted by,", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 10)
    para(doc, "[YOUR NAME] ([USN])", 12.5, True, WD_ALIGN_PARAGRAPH.CENTER, 24)
    para(doc, "Course Name: Data Security and Privacy", 12, True, WD_ALIGN_PARAGRAPH.CENTER, 12)
    para(doc, "Course Code: 22AI73", 12, True, WD_ALIGN_PARAGRAPH.CENTER, 12)
    para(doc, "Sem: VII", 12, True, WD_ALIGN_PARAGRAPH.CENTER, 32)
    para(doc, "Under the Guidance of", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 12)
    para(doc, "[GUIDE NAME]", 12.5, True, WD_ALIGN_PARAGRAPH.CENTER, 9)
    para(doc, "Professor", 12, True, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para(doc, "Department of Artificial Intelligence & Machine Learning", 11.5, True, WD_ALIGN_PARAGRAPH.CENTER, 10)
    para(doc, "DSCE, Bengaluru", 11.5, False, WD_ALIGN_PARAGRAPH.CENTER, 10)
    para(doc, "Academic Year: 2026-27", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 0)


def certificate_page(doc, logos):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec, 0.60, 0.65, 0.75, 0.75)
    set_page_border(sec)
    clear_header_footer(sec)
    para(doc, "Dayananda Sagar College of Engineering", 13.5, True, WD_ALIGN_PARAGRAPH.CENTER, 5, color=BLUE)
    para(doc, "(An Autonomous Institute affiliated to Visvesvaraya Technological University (VTU), Belagavi, Approved by", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para(doc, "AICTE and UGC, Accredited by NAAC with ‘A’ grade)", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 3)
    para(doc, "Shavige Malleshwara Hills, Kumaraswamy Layout, Bengaluru - 560 111. India", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para(doc, "DEPARTMENT OF", 11, True, WD_ALIGN_PARAGRAPH.CENTER, 3, color="0000AA")
    para(doc, "ARTIFICIAL INTELLIGENCE & MACHINE LEARNING", 13.5, True, WD_ALIGN_PARAGRAPH.CENTER, 10, color="0000AA")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(logos[0]), width=Inches(0.78)); p.paragraph_format.space_after = Pt(17)
    para(doc, "CERTIFICATE", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 25, color=BROWN)
    certificate = ("This is to certify that the micro project report entitled “AI Assisted Enterprise Cyber Risk Management Using Large Language Models” carried out by [YOUR NAME] ([USN]) is the bonafide work submitted in partial fulfillment of the requirements of the Data Security and Privacy course for the Bachelor of Engineering programme in Artificial Intelligence & Machine Learning during the academic year 2026-27. The work is based on a critical study of LLM-assisted cybersecurity risk management and has been prepared in accordance with the department requirements.")
    para(doc, certificate, 11.5, False, None, 14, 0, 1.45)
    para(doc, "", space_after=62)
    sig = doc.add_table(rows=2, cols=2); sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in sig.rows[0].cells:
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("------------------------------"); set_run_font(r, 11)
    labels = [("[GUIDE NAME]", "Professor"), ("[HOD NAME]", "Professor & HOD")]
    for cell, (name, role) in zip(sig.rows[1].cells, labels):
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(name + "\n"); set_run_font(r, 11.5, True); r=p.add_run(role); set_run_font(r, 10.5)


def declaration_page(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec, 0.60, 0.65, 0.75, 0.75)
    set_page_border(sec)
    clear_header_footer(sec)
    para(doc, "Dayananda Sagar College of Engineering", 13.5, True, WD_ALIGN_PARAGRAPH.CENTER, 5, color=BLUE)
    para(doc, "(An Autonomous Institute affiliated to Visvesvaraya Technological University (VTU), Belagavi, Approved by", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para(doc, "AICTE and UGC, Accredited by NAAC with ‘A’ grade)", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 3)
    para(doc, "Shavige Malleshwara Hills, Kumaraswamy Layout, Bengaluru - 560 111. India", 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para(doc, "DEPARTMENT OF", 11, True, WD_ALIGN_PARAGRAPH.CENTER, 3, color="0000AA")
    para(doc, "ARTIFICIAL INTELLIGENCE & MACHINE LEARNING", 13.5, True, WD_ALIGN_PARAGRAPH.CENTER, 34, color="0000AA")
    para(doc, "DECLARATION", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 26, color=BROWN)
    statement = ("I, [YOUR NAME] ([USN]), student of the 7th semester B.E. in Artificial Intelligence & Machine Learning, declare that this micro project report entitled “AI Assisted Enterprise Cyber Risk Management Using Large Language Models” is my original academic work. The report presents my understanding and analysis of the cited research paper. All external ideas, data, and published findings used in this report have been acknowledged in the references.")
    para(doc, statement, 11.5, False, None, 14, 0, 1.45)
    para(doc, "", space_after=130)
    para(doc, "Signature of the Student", 11, False, WD_ALIGN_PARAGRAPH.RIGHT, 4)
    para(doc, "[YOUR NAME] - [USN]", 11, False, WD_ALIGN_PARAGRAPH.RIGHT, 0)


def acknowledgement_page(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec, 0.72, 0.72, 0.92, 0.92)
    clear_header_footer(sec)
    para(doc, "ACKNOWLEDGEMENT", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 22)
    texts = [
        "I express my sincere gratitude to the management of Dayananda Sagar College of Engineering for providing an environment that supports learning and project work.",
        "I thank the Head of the Department and the faculty of the Department of Artificial Intelligence & Machine Learning for their guidance and encouragement.",
        "I am especially grateful to my guide, [GUIDE NAME], for valuable direction and support in preparing this report.",
        "I also thank my family and friends for their encouragement throughout the work."
    ]
    for t in texts:
        para(doc, t, 11.5, False, None, 12, 0, 1.45)
    para(doc, "[YOUR NAME] - [USN]", 11, False, WD_ALIGN_PARAGRAPH.RIGHT, 0)


def vision_page(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec, 0.65, 0.65, 0.85, 0.85)
    set_page_border(sec)
    header_footer(sec, page_number=False)
    heading(doc, "Institute Vision", 1)
    para(doc, "To impart quality technical education with a focus on Research and Innovation emphasizing on Development of Sustainable and Inclusive Technology for the benefit of society.", 11.5, space_after=13)
    heading(doc, "Institute Mission", 1)
    for item in ["To provide an environment that enhances creativity and innovation in pursuit of excellence.", "To nurture teamwork in order to transform individuals as responsible leaders and entrepreneurs.", "To train students to the changing technical scenario and make them understand the importance of sustainable and inclusive technologies."]:
        p=para(doc, "•   " + item, 11.5, space_after=7); p.paragraph_format.left_indent=Inches(.20); p.paragraph_format.first_line_indent=Inches(-.20)
    heading(doc, "Department Vision", 1)
    para(doc, "To provide progressive education and flourish the students’ ingenuity to be successful professionals impacting the society for a smarter and ethical world.", 11.5, space_after=13)
    heading(doc, "Department Mission", 1)
    for item in ["To adopt an engaging teaching-learning process with emphasis on problem-solving, programming, and intelligent systems development skills.", "To promote additional skill development through enhanced and experiential learning.", "To collaborate with industries and professional bodies to make students industry-ready.", "To encourage innovation through multi-disciplinary research and development activities.", "To nurture human values and ethical principles in students, empowering them to contribute meaningfully to society as responsible professionals."]:
        p=para(doc, "•   " + item, 11.5, space_after=7); p.paragraph_format.left_indent=Inches(.20); p.paragraph_format.first_line_indent=Inches(-.20)


def toc_pages(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec, 0.65, 0.65, 0.85, 0.85); header_footer(sec, page_number=False)
    para(doc, "Table of Contents", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 18)
    entries = [("1", "Abstract", "1"), ("2", "Introduction", "2"), ("3", "Problem Statement and Objectives", "3"), ("4", "Proposed Solution and Methodology", "4-6"), ("5", "Security and Privacy Aspects", "7"), ("6", "Implementation and Results", "8-10"), ("7", "Testing and Validation", "11"), ("8", "Conclusion and Future Enhancements", "12"), ("9", "References", "13"), ("10", "Appendix - Figure Placeholders", "14-16")]
    table=doc.add_table(rows=1, cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths=[.65,4.9,.75]
    for c,w,t in zip(table.rows[0].cells,widths,["Sl. No.","Contents","Page No."]):
        c.width=Inches(w); set_cell_shading(c,"D9EAF7"); p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(t);set_run_font(r,10.5,True)
    for no,title,page in entries:
        cells=table.add_row().cells
        for c,w,t,a in zip(cells,widths,[no,title,page],[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER]):
            c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];p.alignment=a;p.paragraph_format.space_after=Pt(2);r=p.add_run(t);set_run_font(r,10.5)
    set_table_borders(table)
    sec2=doc.add_section(WD_SECTION.NEW_PAGE); set_margins(sec2,0.65,0.65,0.85,0.85); header_footer(sec2,page_number=False)
    para(doc, "List of Tables", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 18)
    list_table(doc, [("Table 4.1", "Input telemetry groups", "5"), ("Table 6.1", "Reported comparative performance", "9"), ("Table 7.1", "Validation scenarios", "11")])
    para(doc, "", space_after=12)
    para(doc, "List of Figures", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 18)
    list_table(doc, [("Figure 4.1", "Proposed risk analysis workflow placeholder", "4"), ("Figure 4.2", "System architecture placeholder", "6"), ("Figure 6.1", "Dashboard or project interface placeholder", "10"), ("Figure A.1", "Project screenshot placeholder", "14")])


def list_table(doc, rows):
    table=doc.add_table(rows=1, cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[1.15,4.45,.75]
    for c,w,t in zip(table.rows[0].cells,widths,["Table/Figure No.","Title","Page No."]):
        c.width=Inches(w);set_cell_shading(c,"D9EAF7");p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(t);set_run_font(r,10,True)
    for row in rows:
        cells=table.add_row().cells
        for c,w,t,a in zip(cells,widths,row,[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER]):
            c.width=Inches(w);p=c.paragraphs[0];p.alignment=a;r=p.add_run(t);set_run_font(r,10)
    set_table_borders(table)


def chapter_start(doc, number, title):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec,0.65,0.65,0.85,0.85);header_footer(sec,page_number=True)
    if number == 1:
        pg=OxmlElement("w:pgNumType");pg.set(qn("w:start"),"1");sec._sectPr.append(pg)
    para(doc, f"CHAPTER - {number}", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 16)
    heading(doc, title, 1)


def placeholder(doc, label, note, height_lines=7):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0,0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, top={"val":"single","sz":"12","color":"808080"}, bottom={"val":"single","sz":"12","color":"808080"}, left={"val":"single","sz":"12","color":"808080"}, right={"val":"single","sz":"12","color":"808080"})
    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(0)
    r=p.add_run(label+"\n");set_run_font(r,11.5,True,color="595959")
    r=p.add_run(note+"\n" * height_lines);set_run_font(r,10,color="595959",italic=True)
    p.paragraph_format.space_after=Pt(12)
    cap=para(doc, label.replace("PLACEHOLDER", ""), 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 13)
    return cap


def comparison_table(doc):
    para(doc, "Table 6.1 Reported comparative performance", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 5)
    table=doc.add_table(rows=1, cols=4); table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[2.0,1.2,1.2,1.6]
    for c,w,t in zip(table.rows[0].cells,widths,["Model","Accuracy","Macro F1","Notes"]):
        c.width=Inches(w);set_cell_shading(c,"17365D");p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(t);set_run_font(r,9.5,True,"FFFFFF")
    rows=[("Random Forest","0.9421","0.9426","Classical baseline"),("CNN-LSTM","0.9662","0.9663","Strong DL baseline"),("GPT-3.5 Turbo","0.9873","0.9874","API model"),("Mistral-7B-v0.1","0.9936","0.9936","Best reported model")]
    for idx,row in enumerate(rows):
        cells=table.add_row().cells
        for c,w,t in zip(cells,widths,row):
            c.width=Inches(w)
            if idx%2: set_cell_shading(c,"F2F7FB")
            p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER if c!=cells[3] else WD_ALIGN_PARAGRAPH.LEFT;r=p.add_run(t);set_run_font(r,9.5, bold=(row[0].startswith("Mistral")))
    set_table_borders(table)


def body(doc):
    chapter_start(doc, 1, "ABSTRACT")
    para(doc, "This report examines AI-assisted cybersecurity risk management for enterprise environments that generate high-volume, high-velocity network telemetry. It studies the use of instruction-following large language models to classify events such as distributed denial-of-service attacks, intrusions, and malware after structured telemetry is converted into a clear natural-language representation.", 11.5, space_after=12)
    para(doc, "The referenced study evaluates GPT-3.5 Turbo and Mistral-7B against conventional machine-learning and deep-learning baselines. It reports that Mistral-7B achieved the strongest overall classification result, while lighter models remained more resource-efficient. The key practical conclusion is that LLMs should support, rather than replace, a layered Security Operations Center workflow.", 11.5, space_after=12)
    para(doc, "The proposed project view is an AI-assisted risk analyzer that receives telemetry, creates structured prompts, produces a risk category, and displays a traceable explanation with the supporting context. The document highlights deployment trade-offs involving accuracy, latency, compute cost, privacy, and prompt robustness.", 11.5, space_after=8)

    chapter_start(doc, 2, "INTRODUCTION")
    heading(doc, "2.1 Background", 2)
    para(doc, "Digital enterprises rely on cloud services, connected devices, and distributed applications. These systems continuously produce security logs, flow records, alerts, and contextual signals. Their scale makes manual triage difficult, while missed or delayed findings can affect confidentiality, integrity, service availability, regulatory compliance, and stakeholder trust.", 11.5)
    heading(doc, "2.2 Role of AI in Risk Management", 2)
    para(doc, "Traditional controls use signatures, rules, and static thresholds. They remain important, but they can be less effective when threats change rapidly or when several weak signals must be interpreted together. AI-based analytics can learn patterns from telemetry and provide risk scores that help analysts prioritize investigation.", 11.5)
    heading(doc, "2.3 Large Language Models for Security Analytics", 2)
    para(doc, "Instruction-following LLMs can process a human-readable serialization of network data. This gives the analyst a potential bridge between raw event fields and an understandable decision. The value must be evaluated with operational measures such as latency, resource use, calibration, robustness, and the ability to generalize beyond familiar attack patterns.", 11.5)

    chapter_start(doc, 3, "PROBLEM STATEMENT AND OBJECTIVES")
    heading(doc, "3.1 Problem Statement", 2)
    para(doc, "Enterprise SOC teams receive heterogeneous telemetry from network devices, firewalls, intrusion detection systems, user accounts, and endpoints. A classifier may label an event as suspicious, yet analysts still need useful context about the event and confidence in the outcome. A solution that focuses only on accuracy can be impractical if it is slow, costly, opaque, or fragile under changing inputs.", 11.5)
    heading(doc, "3.2 Objectives", 2)
    objectives=["Represent selected security telemetry as a consistent natural-language prompt for LLM-assisted classification.","Compare LLM-based detection with classical and deep-learning baselines using reported accuracy and F1 measures.","Consider deployment factors including latency, memory use, operational cost, and robustness.","Present decisions in an analyst-oriented risk-analysis workflow with a place for project-specific visual evidence."]
    for i,o in enumerate(objectives,1): para(doc,f"{i}. {o}",11.5,space_after=6)
    heading(doc, "3.3 Scope", 2)
    para(doc, "The analysis focuses on network-based attack classes consolidated as DDoS, intrusion, and malware. It discusses a decision-support workflow for a SOC; it does not claim to replace professional incident response, threat hunting, or governance controls.", 11.5)

    chapter_start(doc, 4, "PROPOSED SOLUTION AND METHODOLOGY")
    heading(doc, "4.1 Proposed Workflow", 2)
    para(doc, "The suggested analyzer follows a staged workflow: telemetry ingestion, data validation, prompt construction, LLM or baseline prediction, risk presentation, explanation, and analyst review. Sensitive fields should be minimized or masked before they are passed to an external service.", 11.5)
    placeholder(doc,"FIGURE 4.1 PLACEHOLDER","Insert your project workflow diagram here",5)
    heading(doc, "4.2 Telemetry Inputs", 2)
    para(doc, "The source paper groups inputs into network-flow attributes, security indicators, and contextual logs. Useful examples include source and destination addresses, ports, protocol, packet length, anomaly score, alert severity, firewall output, IDS/IPS alerts, device information, and log source.", 11.5)
    para(doc, "Table 4.1 Input telemetry groups", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 4)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,t in zip(table.rows[0].cells,["Group","Examples"]):set_cell_shading(c,"17365D");p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(t);set_run_font(r,10,True,"FFFFFF")
    for i,row in enumerate([("Network flow","Source/destination, ports, protocol, packet length"),("Security indicators","Malware indicators, anomaly score, alerts, severity"),("Context and logs","User/device metadata, firewall output, IDS/IPS alerts")]):
        cells=table.add_row().cells
        for c,t in zip(cells,row):
            if i%2:set_cell_shading(c,"F2F7FB")
            p=c.paragraphs[0];r=p.add_run(t);set_run_font(r,10)
    set_table_borders(table)
    heading(doc,"4.3 Data Preparation and Prompt Construction",2)
    para(doc,"Categorical fields are normalized or encoded, records are checked for missing values, and the three target classes are consistently represented. Each event can then be serialized into a concise prompt that preserves relationships such as source, destination, protocol, packet length, severity, and anomaly score. Stratified training, validation, and test partitions help preserve class balance during evaluation.",11.5)
    heading(doc,"4.4 Model Options",2)
    para(doc,"The research compares Random Forest and SVM baselines, deep models such as CNN, LSTM, and CNN-LSTM, and instruction-following LLMs. GPT-3.5 Turbo provides an API-based option, whereas Mistral-7B can support local or controlled deployment when adequate compute is available. The model choice should reflect both security performance and the organization’s constraints.",11.5)
    placeholder(doc,"FIGURE 4.2 PLACEHOLDER","Insert your system architecture diagram here",5)

    chapter_start(doc,5,"SECURITY AND PRIVACY ASPECTS")
    heading(doc,"5.1 Data Minimization",2)
    para(doc,"Telemetry can contain identifiers, location details, user metadata, and internal infrastructure information. The analyzer should collect only the fields needed for the defined detection task, mask sensitive values where possible, and enforce access controls for data, prompts, and outputs.",11.5)
    heading(doc,"5.2 Secure Processing",2)
    para(doc,"Raw inputs should be validated before processing. Logs should be protected in transit and at rest, access should be auditable, and retention should follow institutional or organizational policy. When an external LLM endpoint is used, the prompt should avoid unnecessary confidential content and the provider’s data-handling terms must be considered.",11.5)
    heading(doc,"5.3 Model and Prompt Risks",2)
    para(doc,"An LLM-based classifier can be affected by distribution shifts, maliciously crafted inputs, and prompt variation. A robust implementation should test input perturbations, monitor confidence and error patterns, and route ambiguous predictions to a human analyst. The model output is a decision-support signal rather than proof of compromise.",11.5)

    chapter_start(doc,6,"IMPLEMENTATION AND RESULTS")
    heading(doc,"6.1 Reference Implementation Plan",2)
    para(doc,"A practical prototype can use Python for preprocessing and a web interface for the analyst workflow. The interface may show submitted telemetry, a normalized prompt summary, the predicted category, confidence or calibration information, related alerts, and an explanation panel. The final interface and screenshots should be inserted in the supplied figure locations.",11.5)
    heading(doc,"6.2 Reported Results from the Base Paper",2)
    para(doc,"The referenced empirical study reports a performance advantage for the LLM approaches over the listed baselines. Mistral-7B-v0.1 produced the highest reported accuracy and macro-F1 values. This finding must be interpreted with the study conditions in mind; performance can change with a different data source, prompt scheme, model configuration, or operational environment.",11.5)
    comparison_table(doc)
    para(doc,"The study also reports a trade-off: lighter baselines complete inference with less resource use, while LLMs require greater compute. The paper recommends placing LLMs downstream of lightweight filtering or anomaly-detection stages so their semantic reasoning is used where it adds the most value.",11.5,space_after=12)
    placeholder(doc,"FIGURE 6.1 PLACEHOLDER","Insert your dashboard, console, or result screenshot here",5)

    chapter_start(doc,7,"TESTING AND VALIDATION")
    heading(doc,"7.1 Validation Approach",2)
    para(doc,"A project implementation should be validated with representative, permissioned telemetry and clearly separated training, validation, and test partitions. Evaluation should combine classification metrics with operational checks such as inference latency, failure handling, audit logging, and the stability of results when prompt wording or non-critical input order changes.",11.5)
    para(doc,"Table 7.1 Suggested validation scenarios",10.5,True,WD_ALIGN_PARAGRAPH.CENTER,4)
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,t in zip(table.rows[0].cells,["Scenario","Expected behavior","Evidence to attach"]):set_cell_shading(c,"17365D");p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(t);set_run_font(r,9.5,True,"FFFFFF")
    scenarios=[("Valid DDoS-like record","Classifies with risk context","Prediction screenshot"),("Incomplete event","Safely flags or rejects invalid input","Validation message"),("Changed field order","Stable result when semantics are unchanged","Comparison result"),("High-risk finding","Shows escalation guidance, not autonomous remediation","Output screenshot")]
    for i,row in enumerate(scenarios):
        cells=table.add_row().cells
        for c,t in zip(cells,row):
            if i%2:set_cell_shading(c,"F2F7FB")
            p=c.paragraphs[0];r=p.add_run(t);set_run_font(r,9.2)
    set_table_borders(table)
    heading(doc,"7.2 Acceptance Criteria",2)
    para(doc,"The prototype is acceptable when it handles valid and invalid input safely, preserves required event context, produces a readable category and explanation, records sufficient audit information, and does not present the model output as a final security verdict.",11.5)

    chapter_start(doc,8,"CONCLUSION AND FUTURE ENHANCEMENTS")
    heading(doc,"8.1 Conclusion",2)
    para(doc,"LLM-assisted cyber risk management can add contextual interpretation to enterprise security telemetry. The base study reports that Mistral-7B achieved the strongest measured detection performance among the evaluated models, while conventional approaches remained more efficient. The most defensible deployment is therefore a layered workflow in which models support analyst judgment rather than act as an unsupervised replacement for it.",11.5)
    heading(doc,"8.2 Future Enhancements",2)
    for item in ["Add controlled, project-specific screenshots and workflow figures to the marked locations.","Evaluate prompt robustness and adversarial manipulation with approved test data.","Integrate explainability, calibration, and confidence monitoring into the analyst view.","Assess privacy-preserving data handling and controlled local deployment options.","Perform longer-running experiments across diverse enterprise and IoT datasets."]:
        para(doc,"•   "+item,11.5,space_after=6)

    chapter_start(doc,9,"REFERENCES")
    para(doc,"[1] S. Sorour, S. El-Deeb, A. Alojail, and S. A. E. Kahouf, “Security risk management in the digital enterprise: enhancing cyber defense with large language models,” Journal of Big Data, vol. 13, art. 101, 2026. doi: 10.1186/s40537-026-01420-w.",11,space_after=12)
    para(doc,"[2] National Institute of Standards and Technology, Guide for Conducting Risk Assessments, NIST Special Publication 800-30 Revision 1, 2012.",11,space_after=12)
    para(doc,"[3] International Organization for Standardization, ISO/IEC 27005:2022 Information security, cybersecurity and privacy protection - Guidance on managing information security risks, 2022.",11,space_after=12)

    chapter_start(doc,10,"APPENDIX - PROJECT FIGURE PLACEHOLDERS")
    para(doc,"Replace the following boxes with your own approved project figures before submission. Keep each caption with its corresponding figure.",11.5,space_after=12)
    placeholder(doc,"FIGURE A.1 PLACEHOLDER","Insert project home page or data-input screenshot",5)
    placeholder(doc,"FIGURE A.2 PLACEHOLDER","Insert prediction-result screenshot",5)
    placeholder(doc,"FIGURE A.3 PLACEHOLDER","Insert explanation or dashboard screenshot",5)
    placeholder(doc,"FIGURE A.4 PLACEHOLDER","Insert testing or evaluation screenshot",5)


def configure_styles(doc):
    styles=doc.styles
    normal=styles["Normal"]
    normal.font.name="Times New Roman";normal._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman");normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman");normal.font.size=Pt(11.5)
    for name in ["Title","Heading 1","Heading 2"]:
        st=styles[name];st.font.name="Times New Roman";st._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman");st._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman");st.font.color.rgb=RGBColor(0,0,0)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    logos=make_logos()
    doc=Document()
    configure_styles(doc)
    cover_page(doc,logos)
    certificate_page(doc,logos)
    declaration_page(doc)
    acknowledgement_page(doc)
    vision_page(doc)
    toc_pages(doc)
    body(doc)
    doc.core_properties.title="AI Assisted Enterprise Cyber Risk Management Using Large Language Models"
    doc.core_properties.author="[YOUR NAME]"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
